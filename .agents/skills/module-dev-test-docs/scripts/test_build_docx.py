"""Expected: Chinese content and table cells survive; unsupported input fails."""

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from build_docx import build


def self_test():
    with tempfile.TemporaryDirectory(prefix="module-docx-test-") as directory:
        root = Path(directory)
        source, output = root / "case.md", root / "case.docx"
        source.write_text(
            '# 租户管理开发测试文档\n\n中文 **验收结果** 与 `tenant_id`。\n\n'
            '## 测试步骤\n\n- 登录租户甲\n- 查询用户\n\n保留编号 `__platform__`。\n\n'
            '| 字段 | 说明 |\n| --- | --- |\n| `tenant_id` | 保留 A\\|B 与中文 |\n\n'
            '```powershell\nmvn validate\n```\n\n<!-- pagebreak -->\n\n## 结果\n\n待执行。\n',
            encoding="utf-8",
        )
        build(source, output)
        result = Document(output)
        assert result.paragraphs[0].style.name == "Title"
        assert result.paragraphs[0].text == "租户管理开发测试文档"
        assert "中文 验收结果 与 tenant_id。" in [p.text for p in result.paragraphs]
        assert len(result.tables) == 1
        assert result.tables[0].cell(1, 1).text == "保留 A|B 与中文"
        assert result.tables[0].rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
        assert 'w:type="page"' in result.element.xml
        assert any(run.bold and run.text == "验收结果" for paragraph in result.paragraphs for run in paragraph.runs)
        assert any(p.text == "mvn validate" for p in result.paragraphs)
        assert any(p.text == "保留编号 __platform__。" for p in result.paragraphs)
        assert not result.styles['Title'].element.xpath('./w:pPr/w:pBdr')
        assert not result.styles['Title'].element.xpath('./w:rPr/w:rFonts/@w:eastAsiaTheme')
        assert result.paragraphs[0].runs[0].bold is None
        print("PASS Chinese text, formatting, escaped table cells and page breaks")
        for content in ["", "正文缺少标题", "# 标题\n\n![图片](missing.png)", "# 标题\n\n<div>正文</div>"]:
            source.write_text(content, encoding="utf-8")
            try:
                build(source, output)
            except ValueError:
                continue
            raise AssertionError("Unsupported or empty input must fail explicitly")
        print("PASS empty input, missing title, unsupported image and HTML rejection")
        source.write_text("# 标题\n", encoding="utf-8")
        try:
            build(source, source)
        except ValueError:
            pass
        else:
            raise AssertionError("Input overwrite must fail")
        assert source.read_text(encoding="utf-8") == "# 标题\n"
        print("PASS input overwrite protection")
        print("DOCX_BUILDER_TESTS: PASS")


if __name__ == "__main__":
    self_test()
