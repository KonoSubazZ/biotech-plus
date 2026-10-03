"""Convert a constrained Markdown module document to an editable DOCX.

Requires python-docx and markdown-it-py. Unsupported constructs fail explicitly.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from markdown_it import MarkdownIt


def set_font(style, size, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in list(fonts.attrib):
        if attribute.lower().endswith("theme"):
            del fonts.attrib[attribute]
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for border in style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    for name, size, bold in [("Normal", 10, False), ("Title", 22, True),
                             ("Heading 1", 15, True), ("Heading 2", 11.5, True),
                             ("Heading 3", 10.5, True)]:
        style = doc.styles[name]
        set_font(style, size, bold)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.15
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(9)
            style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    for name in ["List Bullet", "List Number"]:
        set_font(doc.styles[name], 9)
        doc.styles[name].paragraph_format.space_after = Pt(4)
        doc.styles[name].paragraph_format.line_spacing = 1.05
    doc.core_properties.title = ""
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    return doc


def add_inline(paragraph, children):
    bold = italic = False
    link = None
    for token in children or []:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "link_open":
            link = token.attrGet("href")
        elif token.type == "link_close":
            if link:
                paragraph.add_run(f" ({link})")
            link = None
        elif token.type in {"text", "code_inline"}:
            run = paragraph.add_run(token.content)
            run.bold = True if bold else None
            run.italic = True if italic else None
            if token.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        else:
            raise ValueError(f"Unsupported inline Markdown: {token.type}")


def table_rows(tokens):
    rows, row = [], []
    for token in tokens:
        if token.type == "tr_open":
            row = []
        elif token.type == "inline":
            row.append(token.children)
        elif token.type == "tr_close":
            rows.append(row)
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Empty or inconsistent table")
    return rows


def add_table(doc, rows):
    count = len(rows[0])
    if count > 5:
        raise ValueError("Use prose or split tables with more than five columns")
    weights = {2: [0.24, 0.76], 3: [0.18, 0.57, 0.25],
               4: [0.18, 0.22, 0.38, 0.22], 5: [0.15, 0.16, 0.24, 0.29, 0.16]}.get(count, [1])
    table = doc.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, weight in enumerate(weights):
        table.columns[index].width = Mm(170 * weight)
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for column_index, children in enumerate(row_data):
            cell = row.cells[column_index]
            cell.width = Mm(170 * weights[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "60" if edge in {"top", "bottom"} else "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell_properties.append(margins)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DFEAF2")
                cell_properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, children)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)


def build(source, output):
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("Input and output must differ")
    if output.suffix.lower() != ".docx":
        raise ValueError("Output must have a .docx suffix")
    text = source.read_text(encoding="utf-8-sig")
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    if not tokens or tokens[0].type != "heading_open" or tokens[0].tag != "h1":
        raise ValueError("Document must start with a level-one title")
    doc, index, list_styles = make_document(), 0, []
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            level = int(token.tag[1:])
            if level == 1 and index != 0:
                raise ValueError("Only one document title is allowed")
            style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, tokens[index + 1].children)
            if index == 0:
                doc.core_properties.title = paragraph.text
            index += 3
        elif token.type == "paragraph_open":
            style = list_styles[-1] if list_styles else "Normal"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, tokens[index + 1].children)
            index += 3
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            list_styles.append("List Bullet" if token.type == "bullet_list_open" else "List Number")
            if len(list_styles) > 1:
                raise ValueError("Use flat lists or headings instead of nested lists")
            index += 1
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_styles.pop()
            index += 1
        elif token.type in {"list_item_open", "list_item_close"}:
            index += 1
        elif token.type == "table_open":
            end = next(i for i in range(index + 1, len(tokens)) if tokens[i].type == "table_close")
            add_table(doc, table_rows(tokens[index:end + 1]))
            index = end + 1
        elif token.type == "fence":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(token.content.rstrip("\n"))
            run.font.name, run.font.size = "Consolas", Pt(8.5)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
            index += 1
        elif token.type == "html_block" and token.content.strip() == "<!-- pagebreak -->":
            doc.add_page_break()
            index += 1
        else:
            raise ValueError(f"Unsupported block Markdown: {token.type}")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    arguments = parser.parse_args()
    document = build(arguments.input, arguments.output)
    print(f"DOCX_CREATED: {arguments.output}")
    print(f"DOCX_STRUCTURE: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables")


if __name__ == "__main__":
    main()
