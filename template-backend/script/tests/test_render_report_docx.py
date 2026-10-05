#!/usr/bin/env python3
"""Unit and integration tests for render_report_docx.py."""

from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

# 本测试放在 template-backend/script/tests/，脚本与模板都在 ruoyi-report 模块里
TEST_FILE = Path(__file__).resolve()
BACKEND_DIR = TEST_FILE.parents[2]
PROJECT_DIR = BACKEND_DIR / "ruoyi-modules/ruoyi-report"
SCRIPT_DIR = PROJECT_DIR / "src/main/resources/scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from render_report_docx import (  # noqa: E402
    InputValidationError,
    TemplateRenderError,
    render_report,
)


class RenderReportDocxTest(unittest.TestCase):

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.work_dir = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def write_template(self, text: str) -> Path:
        path = self.work_dir / "template.docx"
        document = Document()
        document.add_paragraph(text)
        document.save(path)
        return path

    def write_json(self, value: object) -> Path:
        path = self.work_dir / "payload.json"
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def document_text(self, path: Path) -> str:
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_renders_nested_json_and_list(self) -> None:
        template = self.write_template(
            "姓名：{{ sampleInfo.name }}；基因："
            "{% for item in genes %}{{ item.symbol }}{% if not loop.last %}, {% endif %}{% endfor %}"
        )
        payload = self.write_json(
            {"sampleInfo": {"name": "张三"}, "genes": [{"symbol": "BRCA1"}, {"symbol": "BRCA2"}]}
        )
        output = self.work_dir / "output.docx"

        result = render_report(template, payload, output)

        assert result["status"] == "GENERATED"
        assert result["templateVariables"] == ["genes", "sampleInfo"]
        assert self.document_text(output) == "姓名：张三；基因：BRCA1, BRCA2"

    def test_rejects_missing_nested_field_without_writing_output(self) -> None:
        template = self.write_template("{{ sampleInfo.name }}")
        payload = self.write_json({"sampleInfo": {}})
        output = self.work_dir / "output.docx"

        with self.assertRaises(TemplateRenderError):
            render_report(template, payload, output)

        assert not output.exists()

    def test_rejects_non_object_json_root(self) -> None:
        template = self.write_template("固定文本")
        payload = self.write_json([{"name": "错误根节点"}])

        with self.assertRaises(InputValidationError):
            render_report(template, payload, self.work_dir / "output.docx")

    def test_failed_render_keeps_existing_output_unchanged(self) -> None:
        template = self.write_template("{{ required.value }}")
        payload = self.write_json({})
        output = self.work_dir / "output.docx"
        output.write_bytes(b"existing-output")

        with self.assertRaises(TemplateRenderError):
            render_report(template, payload, output)

        assert output.read_bytes() == b"existing-output"

    def test_renders_project_template_from_example_json(self) -> None:
        template = PROJECT_DIR / "src/main/resources/report-templates/同源重组修复（HRR）通路基因检测报告-圣域_v1.docx"
        payload = PROJECT_DIR / "src/main/resources/report-templates/同源重组修复（HRR）通路基因检测报告-圣域_v1.example.json"
        output = self.work_dir / "pharma-shengyu.docx"

        result = render_report(template, payload, output)

        assert result["status"] == "GENERATED"
        with zipfile.ZipFile(output) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        assert "示例研究中心" in document_xml
        assert "乳腺癌" in document_xml
        assert "{{" not in document_xml
        assert "{%" not in document_xml

        preserve_prefixes = (
            "customXml/",
            "word/header",
            "word/footer",
            "word/media/",
            "word/numbering",
            "word/theme/",
        )
        with zipfile.ZipFile(template) as source, zipfile.ZipFile(output) as rendered:
            rendered_names = set(rendered.namelist())
            for name in source.namelist():
                if name.startswith(preserve_prefixes):
                    assert name in rendered_names
                    assert source.read(name) == rendered.read(name)


if __name__ == "__main__":
    unittest.main(verbosity=2)
